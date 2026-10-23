import csv
import io
import zipfile

import phonenumbers
from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

COLUMNS = ("phone_number", "phone")


def _rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        yield from csv.reader(io.TextIOWrapper(upload.file, encoding="utf-8-sig", newline=""))
    elif name.endswith(".xlsx"):
        yield from load_workbook(upload.file, read_only=True, data_only=True).active.iter_rows(values_only=True)
    else:
        raise ValidationError({"contact_file": "Only .csv and .xlsx files are supported."})


def parse_numbers(upload):
    """Return the valid, deduplicated E.164 numbers (first-seen order) from a CSV/XLSX upload."""
    if upload.size > settings.CONTACT_IMPORT_MAX_BYTES:
        raise ValidationError({"contact_file": "File is too large."})
    try:
        rows = iter(_rows(upload))
        header = [str(c or "").strip().lower() for c in next(rows, [])]
        col = next((header.index(c) for c in COLUMNS if c in header), None)
        if col is None:
            raise ValidationError({"contact_file": "A phone_number column is required."})
        numbers = {}
        for row in rows:
            try:
                # No default region: numbers must carry a country code (+...)
                parsed = phonenumbers.parse(str(row[col]).strip(), None)
            except (IndexError, phonenumbers.NumberParseException):
                continue
            if phonenumbers.is_valid_number(parsed):
                numbers.setdefault(phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164))
    except (UnicodeDecodeError, csv.Error, OSError, KeyError, ValueError, TypeError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ValidationError({"contact_file": "The file could not be read."}) from exc
    return list(numbers)
